import io
import os
import openpyxl
from fastapi.testclient import TestClient
from main import app

client = TestClient(app)

def get_test_data():
    if os.path.exists("html1.htm") and os.path.exists("html2.htm"):
        with open("html1.htm", "rb") as f1, open("html2.htm", "rb") as f2:
            return f1.read(), f2.read(), 4, 126, 30, 96

    # Realistic mock HTML files for headless testing
    filler = "<table><tr><td>Header</td></tr></table>\n" * 4
    mock1 = (
        "<html><head><title>Ranking list of Agents</title></head><body>"
        + filler +
        "<table><thead><tr>"
        "<th>Rank</th><th>Branch</th><th>Dev.Code</th><th>Agent</th><th>Agent Name</th>"
        "<th>Phone No.</th><th>SP No</th><th>NSP No</th><th>Tot No</th><th>Sum Ass.</th>"
        "<th>Sing Prem</th><th>Non Sing Prem</th><th>Total Prem</th><th>G&U</th>"
        "</tr></thead><tbody>"
        "<tr><td>1</td><td>488</td><td>D01</td><td>AG001</td><td>John Doe</td><td>1234</td><td>1</td><td>2</td><td>3</td><td>1000</td><td>50</td><td>100</td><td>150</td><td>0</td></tr>"
        "<tr><td>2</td><td>488</td><td>D02</td><td>AG002</td><td>Jane Smith</td><td>5678</td><td>0</td><td>1</td><td>1</td><td>500</td><td>0</td><td>50</td><td>50</td><td>0</td></tr>"
        "</tbody></table></body></html>"
    ).encode("utf-8")

    mock2 = (
        "<html><head><title>Dev.Officer / Agents Statistics</title></head><body>"
        + filler +
        "<table><thead><tr>"
        "<th>Sl.No</th><th>Dev.Code</th><th>Dev.Name</th><th>Agent</th><th>Agent Name</th>"
        "<th>SP POL</th><th>NSP POL</th><th>Total POL</th><th>SING PREM.</th><th>NON SING PREM.</th><th>TOTAL PREM(SP+NSP)</th>"
        "</tr></thead><tbody>"
        "<tr><td>1</td><td>D01</td><td>Manager</td><td>AG001</td><td>John Doe</td><td>1</td><td>2</td><td>3</td><td>50</td><td>100</td><td>150</td></tr>"
        "</tbody></table></body></html>"
    ).encode("utf-8")

    return mock1, mock2, 4, 2, 1, 1

def test_health():
    res = client.get("/health")
    assert res.status_code == 200
    assert res.json()["status"] == "healthy"
    print("Health check passed.")

def test_index_page():
    res = client.get("/")
    assert res.status_code == 200
    assert "MIS HTML to Excel" in res.text
    assert "Made by <strong>Latu Solutions</strong>" in res.text
    assert "currentYear" in res.text
    print("Web UI page passed.")

def test_summary_endpoint():
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()
    res = client.post(
        "/api/summary",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx}
    )
    assert res.status_code == 200
    data = res.json()
    assert data["status"] == "success"
    assert data["summary"]["total_agents"] == total_agents
    assert data["summary"]["matched"] == matched
    assert data["summary"]["unmatched"] == unmatched
    print("Summary API endpoint passed:", data["summary"])

def test_convert_endpoint():
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()
    res = client.post(
        "/api/convert",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx}
    )
    assert res.status_code == 200
    assert res.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert res.headers["x-stats-matched"] == str(matched)
    assert res.headers["x-stats-unmatched"] == str(unmatched)
    assert res.headers["x-stats-total-agents"] == str(total_agents)

    # Verify that the returned bytes are a valid Excel workbook in memory
    content = res.content
    assert len(content) > 1000
    wb = openpyxl.load_workbook(io.BytesIO(content))
    sheets = wb.sheetnames
    assert "HTML1 - Ranking" in sheets
    assert "HTML2 - Period Stats" in sheets
    assert "Merged (VLOOKUP)" in sheets
    assert "Summary" in sheets
    print("Convert endpoint passed! Generated valid Excel workbook with sheets:", sheets)

def test_dynamic_multi_files():
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()
    res = client.post(
        "/api/convert",
        files=[
            ("file1", ("baseReport.html", f1_data, "text/html")),
            ("comparison_files", ("temadata.htm", f2_data, "text/html")),
            ("comparison_files", ("dvnDor.html", f2_data, "text/html")),
        ],
        data={"table_index": tbl_idx}
    )
    assert res.status_code == 200
    assert res.headers["x-stats-matched"] == str(matched)
    assert res.headers["x-stats-files-count"] == "2"

    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    sheets = wb.sheetnames
    assert "HTML1 - Ranking" in sheets
    assert "temadata - Raw" in sheets
    assert "dvnDor - Raw" in sheets
    assert "Merged (VLOOKUP)" in sheets
    assert "Summary" in sheets

    ws = wb["Merged (VLOOKUP)"]
    headers = [str(cell.value) for cell in ws[5]]
    assert "Total POL (temadata)" in headers
    assert "Total Prem (temadata)" in headers
    assert "Total POL (dvnDor)" in headers
    assert "Total Prem (dvnDor)" in headers
    print("Dynamic multi-file test passed with custom filenames:", sheets)

def test_custom_column_indices():
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()
    res = client.post(
        "/api/convert",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx, "nop_col_index": 7, "prem_col_index": 9}
    )
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Merged (VLOOKUP)"]
    headers = [str(cell.value) for cell in ws[4]]
    assert "Total POL (html2)" in headers
    assert "NON SING PREM. (html2)" in headers
    print("Custom column indices test (Python 0-based 7 & 9) passed:", headers)

def test_excel_1based_indices():
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()
    res = client.post(
        "/api/convert",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx, "nop_col_index": 8, "prem_col_index": 11, "index_base": 1}
    )
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Merged (VLOOKUP)"]
    headers = [str(cell.value) for cell in ws[4]]
    assert "Total POL (html2)" in headers
    assert "TOTAL PREM(SP+NSP) (html2)" in headers
    print("Excel 1-based indices test (Col 8 & 11) passed:", headers)

def test_dynamic_output_filename_with_date():
    from datetime import datetime
    today_str = datetime.now().strftime("%Y-%m-%d")
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()

    # 1. Default filename: report_with_lookup_YYYY-MM-DD.xlsx
    res_default = client.post(
        "/api/convert",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx}
    )
    assert res_default.status_code == 200
    expected_default_name = f"report_with_lookup_{today_str}.xlsx"
    assert res_default.headers["x-output-filename"] == expected_default_name
    assert f'filename="{expected_default_name}"' in res_default.headers["content-disposition"]

    # 2. Custom prefix filename: branch_mis_summary_YYYY-MM-DD.xlsx
    res_custom = client.post(
        "/api/convert",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx, "output_filename": "branch_mis_summary"}
    )
    assert res_custom.status_code == 200
    expected_custom_name = f"branch_mis_summary_{today_str}.xlsx"
    assert res_custom.headers["x-output-filename"] == expected_custom_name
    assert f'filename="{expected_custom_name}"' in res_custom.headers["content-disposition"]
    print(f"Dynamic date-stamped output filename test passed! Generated: '{expected_default_name}' and '{expected_custom_name}'")

def test_dynamic_merged_sheet_heading():
    f1_data, f2_data, tbl_idx, total_agents, matched, unmatched = get_test_data()

    # Test with custom heading matching user prompt: "RELSUTS FOR THE DATA"
    custom_heading = "RELSUTS FOR THE DATA"
    res = client.post(
        "/api/convert",
        files={
            "file1": ("html1.htm", f1_data, "text/html"),
            "file2": ("html2.htm", f2_data, "text/html")
        },
        data={"table_index": tbl_idx, "report_heading": custom_heading}
    )
    assert res.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    ws = wb["Merged (VLOOKUP)"]

    # 1. Heading text and location
    assert ws["A1"].value == custom_heading

    # 2. Styling: Big size font (16pt bold white)
    cell_a1 = ws["A1"]
    assert cell_a1.font.size == 16
    assert cell_a1.font.bold is True
    assert cell_a1.font.color.rgb == "00FFFFFF" or cell_a1.font.color.rgb == "FFFFFF"

    # 3. Styling: Executive navy fill and row height
    assert cell_a1.fill.fgColor.rgb == "001F4E79" or cell_a1.fill.fgColor.rgb == "1F4E79"
    assert ws.row_dimensions[1].height == 42

    # 4. Merged dynamically across all data columns
    merged_ranges = [str(r) for r in ws.merged_cells.ranges]
    assert any(r.startswith("A1:") for r in merged_ranges), f"Expected A1 merged range, got: {merged_ranges}"
    row1_merge = [r for r in merged_ranges if r.startswith("A1:")][0]
    print(f"Dynamic stylish heading test passed! Range: {row1_merge}, Heading: '{custom_heading}', Height: 42, Font: 16pt Bold")

if __name__ == "__main__":
    test_health()
    test_index_page()
    test_summary_endpoint()
    test_convert_endpoint()
    test_dynamic_multi_files()
    test_custom_column_indices()
    test_excel_1based_indices()
    test_dynamic_output_filename_with_date()
    test_dynamic_merged_sheet_heading()
    print("\nALL IN-MEMORY, DYNAMIC MULTI-FILE, COLUMN INDEX, DATE FILENAME, AND STYLISH MERGED HEADING TESTS PASSED PERFECTLY!")


